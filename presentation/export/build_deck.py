#!/usr/bin/env python3
"""Build PPTX and PDF from captured 16:9 slide PNGs."""
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Emu

ROOT = Path(__file__).resolve().parents[1]
EXPORT = Path(__file__).resolve().parent
SLIDES = [EXPORT / f"slide-0{i}.png" for i in range(1, 6)]

WIDTH_IN = 13.333333
HEIGHT_IN = 7.5


def build_pptx():
    prs = Presentation()
    prs.slide_width = Inches(WIDTH_IN)
    prs.slide_height = Inches(HEIGHT_IN)
    blank = prs.slide_layouts[6]
    for image in SLIDES:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(
            str(image), 0, 0, prs.slide_width, prs.slide_height
        )
    dest = ROOT / "AI-Incident-Copilot-Airport-Middleware.pptx"
    prs.save(dest)
    print("wrote", dest)


def build_pdf():
    images = [Image.open(p).convert("RGB") for p in SLIDES]
    dest = ROOT / "AI-Incident-Copilot-Airport-Middleware.pdf"
    images[0].save(
        dest,
        save_all=True,
        append_images=images[1:],
        resolution=150,
        quality=95,
    )
    print("wrote", dest)


if __name__ == "__main__":
    missing = [p for p in SLIDES if not p.exists()]
    if missing:
        raise SystemExit(f"Missing slides: {missing}")
    build_pptx()
    build_pdf()

#!/usr/bin/env python3
"""Build the Smart Connect innovation-idea PowerPoint.

This is a discussion deck, not a technical design document.
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
from PIL import Image, ImageEnhance, ImageFilter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from lxml import etree


ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
OUTFILE = ROOT / "Smart-Connect.pptx"

# --- Palette: airport signage, not startup-gradient ---
NAVY = RGBColor(0x0C, 0x22, 0x3F)
NAVY_DEEP = RGBColor(0x08, 0x18, 0x2C)
NAVY_MID = RGBColor(0x16, 0x36, 0x5A)
SAND = RGBColor(0xF3, 0xEE, 0xE4)
CREAM = RGBColor(0xFB, 0xF8, 0xF2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
AMBER = RGBColor(0xC4, 0x86, 0x14)
AMBER_SOFT = RGBColor(0xF4, 0xE6, 0xC4)
TEAL = RGBColor(0x1E, 0x68, 0x65)
TEAL_SOFT = RGBColor(0xDC, 0xEB, 0xEA)
CHARCOAL = RGBColor(0x2A, 0x28, 0x24)
MUTED = RGBColor(0x6B, 0x65, 0x59)
LINE = RGBColor(0xDD, 0xD5, 0xC7)
RISK = RGBColor(0xB4, 0x3A, 0x2E)
RISK_SOFT = RGBColor(0xF6, 0xE6, 0xE2)
INK = RGBColor(0x1C, 0x1B, 0x18)

FONT_SANS = "Arial"
FONT_SERIF = "Georgia"

SLIDE_W = 13.333
SLIDE_H = 7.5


# ---------------------------------------------------------------------------
# Image prep
# ---------------------------------------------------------------------------

def icon_to_transparent(src: Path, dest: Path, tint: tuple[int, int, int] | None = None) -> None:
    """Knock out the sand background and optionally recolor the silhouette."""
    im = Image.open(src).convert("RGBA")
    im.thumbnail((512, 512), Image.Resampling.LANCZOS)
    arr = np.array(im).astype(np.float32)
    r, g, b = arr[:, :, 0], arr[:, :, 1], arr[:, :, 2]
    lum = 0.299 * r + 0.587 * g + 0.114 * b
    tint = tint or (12, 34, 63)
    alpha = np.where(lum >= 210, 0, np.where(lum >= 90, np.clip(255 * (1 - (lum - 90) / 120), 0, 255), 255))
    out = np.zeros_like(arr)
    out[:, :, 0] = tint[0]
    out[:, :, 1] = tint[1]
    out[:, :, 2] = tint[2]
    out[:, :, 3] = alpha
    im = Image.fromarray(out.astype(np.uint8), "RGBA")
    bbox = im.getbbox()
    if bbox:
        im = im.crop(bbox)
        side = int(max(im.size) * 1.18)
        canvas = Image.new("RGBA", (side, side), (0, 0, 0, 0))
        canvas.paste(im, ((side - im.size[0]) // 2, (side - im.size[1]) // 2), im)
        im = canvas
    im.save(dest)


def prepare_assets() -> dict[str, Path]:
    prepared = ASSETS / "_prepared"
    prepared.mkdir(exist_ok=True)
    mapping = {}
    tints = {
        "navy": (12, 34, 63),
        "white": (255, 255, 255),
        "amber": (196, 134, 20),
        "teal": (30, 104, 101),
        "cream": (251, 248, 242),
    }
    names = ["passenger", "airplane", "baggage", "gate", "clock", "staff", "route"]
    for name in names:
        src = ASSETS / f"icon-{name}.png"
        for tint_name, rgb in tints.items():
            dest = prepared / f"{name}-{tint_name}.png"
            if not dest.exists() or dest.stat().st_mtime < src.stat().st_mtime:
                icon_to_transparent(src, dest, rgb)
            mapping[f"{name}-{tint_name}"] = dest

    # Cover crop: right-hand terminal scene, slightly muted.
    cover_src = ASSETS / "title-terminal.jpg"
    if not cover_src.exists():
        cover_src = ASSETS / "title-terminal.png"
    cover_dest = prepared / "cover-right.png"
    im = Image.open(cover_src).convert("RGB")
    w, h = im.size
    # Portrait slice for the right panel (panel is ~5.35 x 7.5).
    target_ratio = 5.35 / 7.5
    crop_w = int(h * target_ratio)
    left = min(w - crop_w, int(w * 0.42))
    crop = im.crop((left, 0, left + crop_w, h))
    crop = ImageEnhance.Color(crop).enhance(0.82)
    crop = ImageEnhance.Contrast(crop).enhance(0.96)
    crop = ImageEnhance.Brightness(crop).enhance(0.90)
    crop = crop.filter(ImageFilter.SMOOTH)
    crop.save(cover_dest, quality=92)
    mapping["cover"] = cover_dest
    return mapping


# ---------------------------------------------------------------------------
# Drawing helpers
# ---------------------------------------------------------------------------

def _disable_shadow(shape) -> None:
    try:
        shape.shadow.inherit = False
    except Exception:
        pass
    spPr = shape._element.spPr
    effect = spPr.find(qn("a:effectLst"))
    if effect is not None:
        spPr.remove(effect)
    etree.SubElement(spPr, qn("a:effectLst"))


def rect(slide, x, y, w, h, fill, *, line=None, line_w=1.0, radius=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius is not None else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(line_w)
    if radius is not None:
        try:
            sh.adjustments[0] = radius
        except Exception:
            pass
    _disable_shadow(sh)
    return sh


def oval(slide, x, y, w, h, fill, *, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1.25)
    _disable_shadow(sh)
    return sh


def set_run(run, *, font, size, color, bold=False, italic=False, tracking=None):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    # Force latin/ea typeface (avoids theme substitution).
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rPr.find(qn(f"a:{tag}"))
        if el is None:
            el = etree.SubElement(rPr, qn(f"a:{tag}"))
        el.set("typeface", font)
    if tracking is not None:
        rPr.set("spc", str(int(tracking)))


def textbox(
    slide,
    x,
    y,
    w,
    h,
    text,
    *,
    font=FONT_SANS,
    size=16,
    color=CHARCOAL,
    bold=False,
    italic=False,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    tracking=None,
    margin=0.0,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.anchor = anchor
    m = Inches(margin)
    tf.margin_left = m
    tf.margin_right = m
    tf.margin_top = m
    tf.margin_bottom = m
    lines = text.split("\n") if text is not None else [""]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        p.line_spacing = 1.08
        run = p.add_run()
        run.text = line
        set_run(run, font=font, size=size, color=color, bold=bold, italic=italic, tracking=tracking)
    return box


def picture(slide, path: Path, x, y, w, h):
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))


def notes(slide, copy: str) -> None:
    slide.notes_slide.notes_text_frame.text = copy.strip()


def light_chrome(slide, number: str, icons: dict[str, Path]) -> None:
    rect(slide, 0, 0, SLIDE_W, SLIDE_H, SAND)
    rect(slide, 0, 0, SLIDE_W, 0.07, AMBER)
    textbox(
        slide, 0.7, 7.12, 8.5, 0.28,
        "SMART CONNECT   ·   Innovation idea for discussion",
        size=10, color=MUTED, tracking=80,
    )
    textbox(
        slide, 11.35, 7.10, 1.35, 0.30, number,
        font=FONT_SERIF, size=14, color=NAVY, align=PP_ALIGN.RIGHT,
    )


def dark_chrome(slide, number: str) -> None:
    rect(slide, 0, 0, SLIDE_W, SLIDE_H, NAVY_DEEP)
    textbox(
        slide, 0.7, 7.12, 8.5, 0.28,
        "SMART CONNECT   ·   Innovation idea for discussion",
        size=10, color=RGBColor(0x9A, 0xB0, 0xC4), tracking=80,
    )
    textbox(
        slide, 11.35, 7.10, 1.35, 0.30, number,
        font=FONT_SERIF, size=14, color=AMBER_SOFT, align=PP_ALIGN.RIGHT,
    )


def kicker(slide, label: str, *, dark=False):
    textbox(
        slide, 0.7, 0.28, 8.0, 0.28, label.upper(),
        size=11, color=(AMBER_SOFT if dark else AMBER), bold=True, tracking=280,
    )


def heading(slide, title: str, *, dark=False, y=0.52, size=30, w=11.8):
    textbox(
        slide, 0.7, y, w, 0.7, title,
        font=FONT_SERIF, size=size, color=(WHITE if dark else NAVY),
    )


def pill(slide, x, y, w, h, label, *, fill=RISK, font_color=WHITE, size=11):
    rect(slide, x, y, w, h, fill, radius=0.5)
    textbox(
        slide, x, y, w, h, label,
        size=size, color=font_color, bold=True, align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE, tracking=120,
    )


# ---------------------------------------------------------------------------
# Slides
# ---------------------------------------------------------------------------

def slide_1(prs, icons):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, SLIDE_W, SLIDE_H, NAVY_DEEP)
    picture(s, icons["cover"], 7.95, 0, 5.39, 7.5)
    rect(s, 0, 0, 8.55, SLIDE_H, NAVY)
    rect(s, 8.55, 0, 0.08, SLIDE_H, AMBER)

    textbox(s, 0.75, 0.42, 7.2, 0.28, "INNOVATION IDEA",
            size=11, color=AMBER, bold=True, tracking=320)
    textbox(s, 0.75, 1.85, 7.4, 0.85, "SMART CONNECT",
            font=FONT_SANS, size=40, color=WHITE, bold=True, tracking=420)
    rect(s, 0.78, 2.78, 2.35, 0.055, AMBER)

    textbox(
        s, 0.75, 3.05, 7.2, 1.15,
        "Helping passengers catch connecting flights\nbefore it’s too late.",
        font=FONT_SERIF, size=22, color=RGBColor(0xF3, 0xEE, 0xE4), italic=True,
    )
    textbox(
        s, 0.75, 4.40, 7.0, 0.55,
        "A simple idea to prevent missed connections at airports.",
        size=15, color=RGBColor(0xC5, 0xD0, 0xDC),
    )

    # Route strip
    rect(s, 0.75, 5.55, 6.7, 0.95, NAVY_MID, radius=0.08)
    picture(s, icons["airplane-cream"], 0.95, 5.72, 0.55, 0.55)
    textbox(s, 1.60, 5.68, 5.5, 0.32, "DEL    →    BOM    →    LHR",
            size=16, color=WHITE, bold=True, tracking=80)
    textbox(s, 1.60, 6.00, 5.5, 0.32, "Delhi          Mumbai          London",
            size=11, color=RGBColor(0xB7, 0xC6, 0xD4))

    textbox(s, 0.75, 7.12, 7.5, 0.26,
            "Airport operations  ·  For internal discussion",
            size=10, color=RGBColor(0x8E, 0xA3, 0xB8), tracking=60)
    textbox(s, 11.4, 7.10, 1.3, 0.28, "01",
            font=FONT_SERIF, size=14, color=AMBER_SOFT, align=PP_ALIGN.RIGHT)
    notes(
        s,
        "This is a simple idea we’re calling Smart Connect. It’s about a familiar "
        "airport problem: passengers who are about to miss a connecting flight, "
        "and the airport finding out too late to help. I’ll walk through the problem, "
        "the idea, and why it might be worth exploring — this is a discussion, not a finished system.",
    )


def slide_2(prs, icons):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    light_chrome(s, "02", icons)
    kicker(s, "The problem")
    heading(s, "When a delay meets a tight connection")

    textbox(
        s, 0.7, 1.28, 12.0, 0.35,
        "A passenger is travelling  Delhi  →  Mumbai  →  London.",
        size=16, color=CHARCOAL,
    )

    # Journey cards
    cities = [
        ("DEL", "Delhi", "First flight", NAVY),
        ("BOM", "Mumbai", "Connection", AMBER),
        ("LHR", "London", "Onward flight", TEAL),
    ]
    x0 = 0.7
    card_w, card_h, gap = 3.15, 1.62, 1.15
    for i, (code, city, role, accent) in enumerate(cities):
        x = x0 + i * (card_w + gap)
        rect(s, x, 1.78, card_w, card_h, CREAM, line=LINE, line_w=0.75, radius=0.08)
        rect(s, x, 1.78, 0.10, card_h, accent)
        textbox(s, x + 0.28, 1.88, 2.7, 0.55, code, font=FONT_SERIF, size=28, color=NAVY)
        textbox(s, x + 0.28, 2.42, 2.7, 0.28, city, size=14, color=CHARCOAL)
        textbox(s, x + 0.28, 2.72, 2.7, 0.28, role.upper(), size=10, color=MUTED, tracking=140, bold=True)
        if i < 2:
            ax = x + card_w + 0.12
            rect(s, ax, 2.50, gap - 0.24, 0.035, LINE)
            oval(s, ax + (gap - 0.24) / 2 - 0.11, 2.42, 0.22, 0.22, accent)

    # Delay callout on first hop
    pill(s, 3.48, 1.40, 1.85, 0.32, "+25 MIN DELAY", fill=RISK, size=10)

    textbox(
        s, 0.7, 3.58, 12.0, 0.32,
        "The Delhi–Mumbai flight is delayed by 25 minutes. Now there is very little time to:",
        size=15, color=CHARCOAL,
    )

    tasks = [
        ("passenger-navy", "Reach the next gate"),
        ("gate-navy", "Complete airport processes"),
        ("airplane-navy", "Board the next flight"),
    ]
    for i, (icon, label) in enumerate(tasks):
        x = 0.7 + i * 4.05
        rect(s, x, 4.02, 3.85, 0.95, CREAM, line=LINE, line_w=0.75, radius=0.08)
        picture(s, icons[icon], x + 0.18, 4.18, 0.62, 0.62)
        textbox(s, x + 0.95, 4.22, 2.7, 0.58, label, size=15, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)

    # Baggage + too-late
    rect(s, 0.7, 5.15, 5.85, 0.95, CREAM, line=LINE, line_w=0.75, radius=0.08)
    picture(s, icons["baggage-navy"], 0.88, 5.30, 0.62, 0.62)
    textbox(
        s, 1.62, 5.28, 4.7, 0.72,
        "Checked baggage may not reach\nthe connecting aircraft either.",
        size=14, color=CHARCOAL, anchor=MSO_ANCHOR.MIDDLE,
    )

    rect(s, 6.70, 5.15, 5.95, 0.95, NAVY, radius=0.08)
    textbox(
        s, 6.95, 5.28, 5.5, 0.72,
        "The passenger may only realize there is\na problem when it is already too late.",
        font=FONT_SERIF, size=15, color=WHITE, italic=True, anchor=MSO_ANCHOR.MIDDLE,
    )
    notes(
        s,
        "Here’s a situation we all recognize. A passenger is flying Delhi to Mumbai to London. "
        "The first flight is delayed by twenty-five minutes. Suddenly there is very little time to reach the next gate, "
        "finish airport processes, and board. The bags might not make it either. Often the passenger only realizes "
        "how tight it is when it is already too late.",
    )


def slide_3(prs, icons):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    light_chrome(s, "03", icons)
    kicker(s, "Our idea")
    heading(s, "Smart Connect")

    rect(s, 0.7, 1.35, 11.95, 1.85, CREAM, line=LINE, line_w=0.75, radius=0.08)
    rect(s, 0.7, 1.35, 0.10, 1.85, TEAL)
    textbox(
        s, 1.05, 1.52, 11.3, 1.5,
        "Smart Connect identifies passengers who may be at risk of missing their connecting flight, and alerts the airport early enough to help them.",
        font=FONT_SERIF, size=22, color=NAVY, italic=True, anchor=MSO_ANCHOR.MIDDLE,
    )

    textbox(s, 0.7, 3.42, 12.0, 0.35, "It looks at simple information such as:", size=15, color=CHARCOAL)

    tags = [
        ("airplane-navy", "Flight delay", AMBER),
        ("clock-navy", "Time remaining\nbefore the next flight", NAVY),
        ("gate-navy", "Connecting gate", TEAL),
        ("route-navy", "Passenger’s\nlocation / route", NAVY_MID),
        ("baggage-navy", "Baggage status", RISK),
    ]
    tw, gap = 2.22, 0.18
    x0 = 0.7
    for i, (icon, label, stripe) in enumerate(tags):
        x = x0 + i * (tw + gap)
        rect(s, x, 3.90, tw, 2.15, CREAM, line=LINE, line_w=0.75, radius=0.1)
        rect(s, x, 3.90, tw, 0.10, stripe)
        oval(s, x + 0.16, 4.14, 0.22, 0.22, stripe)  # luggage-tag punch
        picture(s, icons[icon], x + 0.70, 4.18, 0.78, 0.78)
        textbox(
            s, x + 0.12, 5.05, tw - 0.24, 0.80, label,
            size=13, color=NAVY, align=PP_ALIGN.CENTER, bold=True,
        )
    notes(
        s,
        "Smart Connect is a way to spot those passengers earlier. It looks at straightforward information "
        "operations already care about: the delay, how much time is left, where the next gate is, the passenger’s "
        "route through the airport, and whether the bags are moving with them. The point is to alert the airport "
        "while there is still a chance to help. We’re not describing a build — just the idea.",
    )


def slide_4(prs, icons):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    light_chrome(s, "04", icons)
    kicker(s, "How it works")
    heading(s, "Five simple steps")

    steps = [
        ("1", "airplane-navy", "Flight gets delayed", CREAM, NAVY),
        ("2", "passenger-navy", "Smart Connect checks the passenger’s connection", CREAM, NAVY),
        ("3", "clock-navy", "It checks the time, route and baggage", CREAM, NAVY),
        ("4", "baggage-navy", "It identifies HIGH RISK passengers", RISK_SOFT, RISK),
        ("5", "staff-navy", "Airport staff can take action early", TEAL_SOFT, TEAL),
    ]
    y = 1.38
    for i, (num, icon, label, bg, accent) in enumerate(steps):
        # connector
        if i < len(steps) - 1:
            rect(s, 1.07, y + 0.72, 0.045, 0.42, LINE)
        oval(s, 0.78, y + 0.12, 0.62, 0.62, accent)
        textbox(
            s, 0.78, y + 0.12, 0.62, 0.62, num,
            size=18, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
        )
        border = LINE if bg == CREAM else None
        rect(s, 1.62, y, 10.95, 0.88, bg, line=border, line_w=0.75, radius=0.08)
        picture(s, icons[icon], 1.82, y + 0.14, 0.60, 0.60)
        textbox(
            s, 2.60, y, 7.4, 0.88, label,
            size=18, color=NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE,
        )
        if num == "4":
            pill(s, 10.05, y + 0.24, 2.22, 0.40, "HIGH RISK", fill=RISK, size=12)
        if num == "5":
            textbox(
                s, 10.05, y, 2.3, 0.88, "While there\nis still time",
                size=12, color=TEAL, italic=True, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE,
            )
        y += 1.05
    notes(
        s,
        "The flow is simple. A flight is delayed. Smart Connect looks at who has a connection. "
        "It checks time, route, and baggage. It flags high-risk passengers. Then airport staff can step in earlier — "
        "not after the connection is already missed.",
    )


def slide_5(prs, icons):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    light_chrome(s, "05", icons)
    kicker(s, "A simple example")
    heading(s, "Same passenger. Now with time still left.")

    # Boarding-pass card
    rect(s, 0.7, 1.38, 6.55, 5.35, CREAM, line=LINE, line_w=0.9, radius=0.06)
    rect(s, 0.7, 1.38, 6.55, 0.58, NAVY)
    textbox(s, 0.92, 1.40, 3.6, 0.52, "BOARDING PASS",
            size=12, color=AMBER, bold=True, tracking=240, anchor=MSO_ANCHOR.MIDDLE)
    pill(s, 4.85, 1.48, 2.15, 0.38, "HIGH RISK", fill=RISK, size=11)

    textbox(s, 0.95, 2.15, 5.9, 0.22, "PASSENGER JOURNEY", size=10, color=MUTED, tracking=160, bold=True)
    textbox(s, 0.95, 2.38, 5.9, 0.55, "DEL  →  BOM  →  LHR", font=FONT_SERIF, size=26, color=NAVY)
    textbox(s, 0.95, 2.92, 5.9, 0.28, "Delhi          Mumbai          London", size=12, color=MUTED)

    # perforation
    for i in range(14):
        oval(s, 3.88, 3.38 + i * 0.22, 0.11, 0.11, SAND)

    rect(s, 0.95, 3.42, 2.70, 1.55, SAND, radius=0.08)
    textbox(s, 1.08, 3.52, 2.45, 0.28, "FIRST FLIGHT", size=10, color=MUTED, tracking=120, bold=True)
    textbox(s, 1.08, 3.82, 2.45, 0.90, "25-minute\ndelay", font=FONT_SERIF, size=20, color=RISK)

    rect(s, 4.22, 3.42, 2.78, 1.55, SAND, radius=0.08)
    textbox(s, 4.35, 3.52, 2.5, 0.28, "NEXT FLIGHT", size=10, color=MUTED, tracking=120, bold=True)
    textbox(s, 4.35, 3.82, 2.5, 0.70, "45 minutes\nuntil boarding closes", font=FONT_SERIF, size=16, color=NAVY)

    rect(s, 0.95, 5.15, 6.05, 1.32, NAVY, radius=0.08)
    picture(s, icons["clock-cream"], 1.15, 5.42, 0.72, 0.72)
    textbox(s, 2.00, 5.28, 4.7, 0.38, "Smart Connect view", size=11, color=AMBER, tracking=80, bold=True)
    textbox(s, 2.00, 5.58, 4.7, 0.70, "This connection is HIGH RISK.\nThere is still a window to help.",
            size=14, color=WHITE)

    # Actions
    textbox(s, 7.50, 1.38, 5.15, 0.30, "EXAMPLES OF POSSIBLE ACTIONS",
            size=11, color=AMBER, bold=True, tracking=160)
    textbox(
        s, 7.50, 1.68, 5.15, 0.55,
        "These are options the airport could take — not automatic or guaranteed actions.",
        size=12, color=MUTED, italic=True,
    )

    actions = [
        ("route-navy", "Guide the passenger through the fastest available route"),
        ("gate-navy", "Inform the connecting gate"),
        ("staff-navy", "Alert relevant airport staff"),
        ("baggage-navy", "Check or prioritize baggage where operationally possible"),
    ]
    y = 2.35
    for icon, label in actions:
        rect(s, 7.50, y, 5.15, 0.95, CREAM, line=LINE, line_w=0.75, radius=0.08)
        picture(s, icons[icon], 7.68, y + 0.18, 0.58, 0.58)
        textbox(s, 8.40, y, 4.05, 0.95, label, size=14, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
        y += 1.05
    notes(
        s,
        "Same passenger: Delhi–Mumbai–London. Twenty-five minutes late on the first flight, forty-five minutes until boarding closes. "
        "Smart Connect would call this high risk. From there, the airport could choose to help — for example, guiding the passenger "
        "on a faster route, informing the next gate, alerting staff, or checking the bags. These are possible actions, not automatic promises.",
    )


def slide_6(prs, icons):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    light_chrome(s, "06", icons)
    kicker(s, "What makes the idea different?")
    heading(s, "Help before the connection is missed")

    # Traditional
    rect(s, 0.7, 1.40, 5.85, 2.85, CREAM, line=LINE, line_w=0.75, radius=0.08)
    textbox(s, 0.95, 1.55, 5.4, 0.28, "TRADITIONAL APPROACH", size=11, color=MUTED, bold=True, tracking=180)
    textbox(
        s, 0.95, 1.95, 5.4, 1.15,
        "Passenger misses the connection.\nThe problem is handled afterwards.",
        font=FONT_SERIF, size=20, color=CHARCOAL,
    )
    rect(s, 0.95, 3.30, 5.35, 0.08, LINE)
    textbox(s, 0.95, 3.48, 5.4, 0.50, "React when it is already too late.", size=14, color=MUTED, italic=True)

    # Smart Connect
    rect(s, 6.80, 1.40, 5.85, 2.85, NAVY, radius=0.08)
    textbox(s, 7.05, 1.55, 5.4, 0.28, "SMART CONNECT", size=11, color=AMBER, bold=True, tracking=180)
    textbox(
        s, 7.05, 1.95, 5.4, 1.15,
        "Risk is identified earlier.\nThe airport has a chance to help before the connection is missed.",
        font=FONT_SERIF, size=18, color=WHITE,
    )
    rect(s, 7.05, 3.30, 5.35, 0.08, RGBColor(0x2A, 0x4A, 0x6E))
    textbox(s, 7.05, 3.48, 5.4, 0.50, "Act while there is still time.", size=14, color=AMBER_SOFT, italic=True)

    # Key idea
    rect(s, 0.7, 4.45, 7.55, 2.05, CREAM, line=LINE, line_w=0.75, radius=0.08)
    textbox(s, 0.95, 4.58, 7.1, 0.28, "ONE THING TO REMEMBER", size=11, color=AMBER, bold=True, tracking=160)
    textbox(
        s, 0.95, 4.95, 7.1, 1.25,
        "Don’t wait for the passenger to miss the flight.\nIdentify the risk early.",
        font=FONT_SERIF, size=20, color=NAVY, italic=True,
    )

    rect(s, 8.45, 4.45, 4.20, 2.05, TEAL, radius=0.08)
    textbox(s, 8.70, 4.58, 3.75, 0.28, "AND KEEP BOTH IN VIEW", size=11, color=AMBER_SOFT, bold=True, tracking=120)
    picture(s, icons["passenger-cream"], 8.70, 5.00, 0.55, 0.55)
    picture(s, icons["baggage-cream"], 9.35, 5.00, 0.55, 0.55)
    textbox(
        s, 8.70, 5.58, 3.75, 0.70,
        "Passenger + baggage\nSometimes the person makes it. The bag does not.",
        size=13, color=WHITE,
    )
    notes(
        s,
        "Today we often handle missed connections after they happen. This idea is about seeing the risk earlier, "
        "while help is still possible. One line to remember: don’t wait for the passenger to miss the flight. "
        "And it is not only the person — sometimes the passenger makes it and the bag doesn’t. Smart Connect keeps both in view.",
    )


def slide_7(prs, icons):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    light_chrome(s, "07", icons)
    kicker(s, "Benefits")
    heading(s, "Who this could help")

    # Passengers
    rect(s, 0.7, 1.40, 5.85, 4.85, CREAM, line=LINE, line_w=0.75, radius=0.08)
    picture(s, icons["passenger-navy"], 0.95, 1.60, 0.62, 0.62)
    textbox(s, 1.70, 1.68, 4.5, 0.50, "For passengers", font=FONT_SERIF, size=24, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
    pax = [
        ("Less stress", "Fewer last-minute surprises at the gate."),
        ("Better guidance", "Clearer help when time is short."),
        ("Fewer missed connections", "A better chance of making the next flight."),
    ]
    y = 2.50
    for title, sub in pax:
        oval(s, 1.05, y + 0.12, 0.18, 0.18, AMBER)
        textbox(s, 1.45, y, 4.7, 0.32, title, size=16, color=NAVY, bold=True)
        textbox(s, 1.45, y + 0.32, 4.7, 0.40, sub, size=13, color=MUTED)
        y += 1.05

    # Airlines / airports
    rect(s, 6.80, 1.40, 5.85, 4.85, NAVY, radius=0.08)
    picture(s, icons["staff-cream"], 7.05, 1.60, 0.62, 0.62)
    textbox(s, 7.80, 1.68, 4.5, 0.50, "For airlines and airports", font=FONT_SERIF, size=22, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    ops = [
        ("Better coordination", "Teams see the same at-risk connection sooner."),
        ("Earlier intervention", "Time to help, not only to apologize."),
        ("Better passenger experience", "Calmer handling of a stressful moment."),
        ("Fewer knock-on problems", "Potentially fewer missed bags and missed flights."),
    ]
    y = 2.40
    for title, sub in ops:
        oval(s, 7.15, y + 0.10, 0.16, 0.16, AMBER)
        textbox(s, 7.50, y, 4.8, 0.28, title, size=15, color=WHITE, bold=True)
        textbox(s, 7.50, y + 0.28, 4.8, 0.32, sub, size=12, color=RGBColor(0xC5, 0xD0, 0xDC))
        y += 0.88

    notes(
        s,
        "For passengers, this could mean less stress, clearer guidance, and fewer missed connections. "
        "For airlines and airports, it means better coordination and a chance to intervene earlier. "
        "We’re not attaching numbers to this yet. The value is a calmer, more timely response when a connection is at risk.",
    )


def slide_8(prs, icons):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    dark_chrome(s, "08")
    kicker(s, "A possible path", dark=True)
    heading(s, "Start small", dark=True)

    rect(s, 0.7, 1.40, 12.0, 1.15, NAVY_MID, radius=0.08)
    rect(s, 0.7, 1.40, 0.10, 1.15, AMBER)
    textbox(s, 0.95, 1.40, 11.5, 1.15,
            "First, identify high-risk connecting passengers.",
            font=FONT_SERIF, size=24, color=WHITE, italic=True, anchor=MSO_ANCHOR.MIDDLE)

    textbox(s, 0.7, 2.75, 12.0, 0.35, "Later it could potentially expand to:", size=14, color=RGBColor(0xB7, 0xC6, 0xD4))

    later = [
        ("baggage-cream", "Baggage\ncoordination"),
        ("gate-cream", "Gate\ncoordination"),
        ("route-cream", "Passenger\nguidance"),
        ("staff-cream", "Disruption\nassistance"),
    ]
    for i, (icon, label) in enumerate(later):
        x = 0.7 + i * 3.1
        rect(s, x, 3.20, 2.90, 1.55, NAVY_MID, radius=0.08)
        picture(s, icons[icon], x + 0.20, 3.40, 0.52, 0.52)
        textbox(s, x + 0.82, 3.32, 1.90, 1.25, label, size=15, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)

    rect(s, 0.7, 5.05, 12.0, 1.70, RGBColor(0x0A, 0x1C, 0x33), radius=0.08)
    rect(s, 0.7, 5.05, 0.10, 1.70, AMBER)
    textbox(
        s, 1.05, 5.20, 11.3, 1.40,
        "From reacting to missed connections\nto preventing them.",
        font=FONT_SERIF, size=26, color=WHITE, italic=True, anchor=MSO_ANCHOR.MIDDLE,
    )
    notes(
        s,
        "We’d start small: just identify high-risk connecting passengers. Over time, it could grow into baggage and gate "
        "coordination, passenger guidance, or help during wider disruptions. The shift we’re proposing is simple — "
        "from reacting to missed connections, to preventing them.",
    )


def main() -> None:
    icons = prepare_assets()
    prs = Presentation()
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)

    cp = prs.core_properties
    cp.author = "Airport Innovation Team"
    cp.last_modified_by = "Airport Innovation Team"
    cp.title = "Smart Connect — Innovation Idea"
    cp.subject = "Helping passengers catch connecting flights before it’s too late"
    cp.category = "Innovation discussion"
    cp.comments = "Discussion draft. Not a description of a live system."

    slide_1(prs, icons)
    slide_2(prs, icons)
    slide_3(prs, icons)
    slide_4(prs, icons)
    slide_5(prs, icons)
    slide_6(prs, icons)
    slide_7(prs, icons)
    slide_8(prs, icons)

    prs.save(OUTFILE)
    print(f"Wrote {OUTFILE}")


if __name__ == "__main__":
    main()
